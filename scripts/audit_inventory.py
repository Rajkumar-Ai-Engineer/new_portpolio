import openpyxl
import json
import os

os.makedirs("docs/audit", exist_ok=True)
wb = openpyxl.load_workbook("resources/project_details_file/main-inventory.xlsx", data_only=True)
ws = wb["Projects"]

headers = [ws.cell(4, c).value for c in range(1, ws.max_column + 1) if ws.cell(4, c).value is not None]
print("Headers:", headers)

projects = []
for r in range(5, ws.max_row + 1):
    pid = ws.cell(r, 1).value
    title = ws.cell(r, 2).value
    if not title and not pid:
        continue
    row_data = {}
    for c, h in enumerate(headers, 1):
        val = ws.cell(r, c).value
        # format numbers/floats
        if isinstance(val, float) and val.is_integer():
            val = int(val)
        row_data[h] = val
    row_data["_row_number"] = r
    projects.append(row_data)

print(f"Total projects found: {len(projects)}")
for p in projects:
    print(f"Row {p['_row_number']} | {p.get('ID')}: {p.get('Project title')} | Tech: {p.get('Technology')} | Tier: {p.get('Tier')} | Status: {p.get('Status')} | Year: {p.get('Year')}")

with open("docs/audit/raw_projects_inventory.json", "w", encoding="utf-8") as f:
    json.dump(projects, f, indent=2, ensure_ascii=False)

print("Saved docs/audit/raw_projects_inventory.json")
